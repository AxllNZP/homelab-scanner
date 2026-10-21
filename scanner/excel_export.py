"""
excel_export.py
----------------
Exporta un ScanSession a un archivo Excel (.xlsx) -- RF-9, la
mejora que quedó anotada como "Should have" en el SRS original.

Fíjate en algo importante: este archivo es casi un espejo de
report.py en su estructura (recibe un ScanSession, no sabe nada de
Nmap ni de subprocess), pero NO tocamos report.py para nada. Esa es
la prueba de que el diseño modular que armamos desde el principio
(RF-7: separar los datos de la presentación) funcionó como se
esperaba: agregar un formato de salida nuevo fue escribir un
archivo nuevo, no reescribir uno existente.

A diferencia de un PDF (pensado para LEER), un Excel está pensado
para que alguien pueda filtrar, ordenar o cruzar los datos después
-- por eso aquí no nos preocupamos tanto por que "se vea bonito" y
sí por que los datos queden en una tabla limpia, una fila por
puerto, fácil de manipular.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from scanner.models import ScanSession


def generate_excel(sesion: ScanSession, output_path: str = "reporte_scan.xlsx") -> str:
    """
    Genera un archivo Excel a partir de un ScanSession y lo guarda
    en output_path. Devuelve la ruta del archivo generado.

    Un Workbook de openpyxl es como "el archivo completo", y dentro
    puede tener varias "hojas" (sheets) -- usamos dos:
        1. "Metadatos"  -> la misma info de trazabilidad del RF-10
                           (Scan ID, fecha, rango, hosts), igual
                           que en el PDF.
        2. "Resultados" -> una fila por cada puerto encontrado, en
                           TODOS los hosts, con el host al que
                           pertenece. Esta es la hoja pensada para
                           que alguien filtre o cruce datos después.
    """
    libro = Workbook()

    # --- Hoja 1: Metadatos (RF-10) ---
    hoja_meta = libro.active  # la hoja que ya viene creada por defecto
    hoja_meta.title = "Metadatos"

    filas_metadatos = [
        ("Scan ID", sesion.scan_id),
        ("Fecha y hora", sesion.timestamp),
        ("Rango analizado", sesion.target_range),
        ("Hosts detectados", sesion.host_count),
    ]
    for fila in filas_metadatos:
        hoja_meta.append(fila)

    # Ensanchamos un poco las columnas para que el texto no quede
    # cortado al abrir el archivo (openpyxl no lo hace automático).
    hoja_meta.column_dimensions["A"].width = 20
    hoja_meta.column_dimensions["B"].width = 50

    # --- Hoja 2: Resultados, una fila por puerto ---
    hoja_resultados = libro.create_sheet(title="Resultados")

    encabezados = ["Host", "Puerto", "Estado", "Servicio"]
    hoja_resultados.append(encabezados)

    # Le damos estilo a la fila de encabezados -- esto es puramente
    # visual, no afecta los datos. Es el equivalente en openpyxl al
    # TableStyle que usamos en report.py con reportlab.
    relleno_encabezado = PatternFill(
        start_color="2B2B2B", end_color="2B2B2B", fill_type="solid"
    )
    fuente_encabezado = Font(color="FFFFFF", bold=True)
    for columna in range(1, len(encabezados) + 1):
        celda = hoja_resultados.cell(row=1, column=columna)
        celda.fill = relleno_encabezado
        celda.font = fuente_encabezado

    # Recorremos cada host y cada puerto, y agregamos una fila por
    # cada uno -- esto es justo lo que hace que la hoja sea "fácil
    # de filtrar": cada fila es un dato independiente y completo,
    # no hay información agrupada que haya que "desarmar" después.
    for host in sesion.hosts:
        if not host.ports:
            hoja_resultados.append([host.ip, "-", "-", "sin puertos abiertos"])
            continue
        for puerto in host.ports:
            hoja_resultados.append(
                [host.ip, puerto.port, puerto.state, puerto.service]
            )

    # Ajustamos el ancho de las columnas de resultados también.
    anchos = [18, 10, 12, 25]
    for indice, ancho in enumerate(anchos, start=1):
        letra_columna = get_column_letter(indice)
        hoja_resultados.column_dimensions[letra_columna].width = ancho

    libro.save(output_path)
    return output_path
